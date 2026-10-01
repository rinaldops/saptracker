"""Exporta a tabela completa da transação YSREGRASPEP para XLSX.

Uso:
    python scripts/export_ysregraspep.py
    python scripts/export_ysregraspep.py --connection 1 --session 0 --out regras.xlsx

A captura é feita pelo mesmo ``inspecionar_completo`` usado no diagnóstico do
SAPTRACKER; portanto, ``GuiTableControl`` é lido página a página, sem depender
de botão de exportação na transação.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.core.sap_connection import SapConnection  # noqa: E402
from src.core.shell_handlers.table_control import GuiTableControlHandler  # noqa: E402

DEFAULT_TABLE_ID = (
    "/app/con[0]/ses[0]/wnd[0]/usr/"
    "tblSAPMYS_MANUT_REGRAS_PEPTC_REGRAS"
)
DEFAULT_OUTPUT = Path("ysregraspep.xlsx")


def _unique_headers(columns: list[str], titles: list[str]) -> list[str]:
    """Monta cabeçalhos Excel não vazios e sem duplicidade."""
    result: list[str] = []
    used: dict[str, int] = {}
    for index, column in enumerate(columns):
        base = (titles[index] if index < len(titles) and titles[index] else column).strip()
        base = base or column or f"Coluna {index + 1}"
        used[base] = used.get(base, 0) + 1
        result.append(base if used[base] == 1 else f"{base} ({used[base]})")
    return result


def export_xlsx(data: dict[str, Any], output: Path) -> int:
    """Grava o resultado da captura como uma planilha XLSX."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError(
            "A dependência openpyxl é necessária. Instale-a com: "
            "python -m pip install openpyxl"
        ) from exc

    columns = [str(column) for column in data.get("colunas", [])]
    titles = [str(title) for title in data.get("titulos", [])]
    headers = _unique_headers(columns, titles)
    rows = data.get("linhas", [])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "YSREGRASPEP"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        sheet.append([row.get(column, "") for column in columns])

    sheet.freeze_panes = "A2"
    if headers:
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            values = [str(sheet.cell(row, index).value or "") for row in range(1, sheet.max_row + 1)]
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(
                60, max(12, max(len(value) for value in values) + 2)
            )

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return len(rows)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--connection", type=int, default=0, help="Índice da conexão SAP (padrão: 0).")
    parser.add_argument("--session", type=int, default=0, help="Índice da sessão SAP (padrão: 0).")
    parser.add_argument("--table-id", default=DEFAULT_TABLE_ID, help="ID completo da GuiTableControl.")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT, help="Arquivo XLSX de saída.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    table_id = args.table_id
    if table_id == DEFAULT_TABLE_ID:
        table_id = table_id.replace('/con[0]/ses[0]/', f'/con[{args.connection}]/ses[{args.session}]/')
    try:
        session = SapConnection().get_session(args.connection, args.session)
        table = session.FindById(table_id)
        if table is None:
            raise RuntimeError(f"Tabela não encontrada: {table_id}")
        data = GuiTableControlHandler().inspecionar_completo(table)
        captured = export_xlsx(data, args.out)
    except Exception as exc:  # noqa: BLE001 - fronteira SAP/arquivo para CLI
        print(f"Erro: {exc}", file=sys.stderr)
        return 2

    print(
        f"Exportado: {args.out} | linhas={captured} | "
        f"esperadas={data.get('total', '?')} | "
        f"páginas={len(data.get('paginas', []))} | "
        f"leituras={data.get('linhas_acumuladas', '?')}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
